"""文档加载器：支持 PDF / Word / TXT / Markdown / CSV / Excel / PPT，输出 LangChain Document。

依赖说明：
- PDF  -> pypdf
- DOCX -> python-docx
- XLSX -> openpyxl（未安装时提示）
- PPTX -> python-pptx（未安装时提示）
- CSV  -> 标准库 csv
- 老版 .doc 不支持，请另存为 .docx 后上传
"""
from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document

from app.exceptions import FileParseError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".pptx"}


def load_file(file_path: str | Path) -> list[Document]:
    """根据扩展名分派解析器，返回 Document 列表（PDF 按页拆分）。"""
    path = Path(file_path)
    if not path.exists():
        raise FileParseError(f"文件不存在: {path.name}")
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise FileParseError(
            f"不支持的文件格式: {ext}，仅支持 {sorted(SUPPORTED_EXTENSIONS)}（老版 .doc 请另存为 .docx）"
        )

    if ext == ".pdf":
        documents = _load_pdf(path)
    elif ext == ".docx":
        documents = _load_docx(path)
    elif ext == ".csv":
        documents = _load_csv(path)
    elif ext == ".xlsx":
        documents = _load_xlsx(path)
    elif ext == ".pptx":
        documents = _load_pptx(path)
    else:
        documents = _load_text(path)

    if not any(doc.page_content.strip() for doc in documents):
        raise FileParseError(f"文件内容为空或无法提取文本: {path.name}")
    return documents


def _base_metadata(path: Path, **extra) -> dict:
    return {"source": str(path), "filename": path.name, **extra}


def _load_text(path: Path) -> list[Document]:
    try:
        content = path.read_text(encoding="utf-8", errors="ignore")
    except OSError as exc:
        raise FileParseError(f"读取文本文件失败: {exc}") from exc
    return [Document(page_content=content, metadata=_base_metadata(path))]


def _load_pdf(path: Path) -> list[Document]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise FileParseError("缺少 pypdf 依赖，请先 pip install pypdf") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # pypdf 异常类型较多，统一包装
        raise FileParseError(f"PDF 解析失败: {exc}") from exc

    documents: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            documents.append(
                Document(page_content=text, metadata=_base_metadata(path, page=index))
            )
    return documents


def _load_docx(path: Path) -> list[Document]:
    try:
        import docx
    except ImportError as exc:
        raise FileParseError("缺少 python-docx 依赖，请先 pip install python-docx") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise FileParseError(f"Word 解析失败: {exc}") from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # 同时提取表格文本，避免漏掉 Word 表格中的内容
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return [Document(page_content="\n".join(paragraphs), metadata=_base_metadata(path))]


def _load_csv(path: Path) -> list[Document]:
    import csv as _csv

    try:
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as f:
            reader = _csv.reader(f)
            rows = [row for row in reader if any(cell.strip() for cell in row)]
    except OSError as exc:
        raise FileParseError(f"读取 CSV 失败: {exc}") from exc
    if not rows:
        raise FileParseError(f"CSV 文件为空: {path.name}")
    # 表头行 + 数据行：每行以竖线分隔，便于检索与展示
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        lines.append(" | ".join(cell.strip() for cell in row))
    return [Document(page_content="\n".join(lines), metadata=_base_metadata(path))]


def _load_xlsx(path: Path) -> list[Document]:
    try:
        import openpyxl
    except ImportError as exc:
        raise FileParseError("缺少 openpyxl 依赖，请先 pip install openpyxl") from exc

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise FileParseError(f"Excel 解析失败: {exc}") from exc

    documents: list[Document] = []
    try:
        for sheet in wb.worksheets:
            lines: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            if lines:
                documents.append(
                    Document(
                        page_content="\n".join(lines),
                        metadata=_base_metadata(path, sheet=sheet.title),
                    )
                )
    finally:
        wb.close()
    return documents


def _load_pptx(path: Path) -> list[Document]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise FileParseError("缺少 python-pptx 依赖，请先 pip install python-pptx") from exc

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise FileParseError(f"PPT 解析失败: {exc}") from exc

    documents: list[Document] = []
    for index, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if lines:
            documents.append(
                Document(
                    page_content="\n".join(lines),
                    metadata=_base_metadata(path, page=index),
                )
            )
    return documents
